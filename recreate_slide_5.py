import os
from PIL import Image, ImageDraw, ImageFont, ImageFilter
from pptx import Presentation
from pptx.util import Inches, Pt
import pymupdf

def create_rounded_rect_mask(size, radius):
    mask = Image.new('L', size, 0)
    draw = ImageDraw.Draw(mask)
    draw.rounded_rectangle((0, 0, size[0], size[1]), radius=radius, fill=255)
    return mask

def draw_hud_corners(draw, rect, color, length=20, width=2):
    x1, y1, x2, y2 = rect
    draw.line([(x1, y1), (x1 + length, y1)], fill=color, width=width)
    draw.line([(x1, y1), (x1, y1 + length)], fill=color, width=width)
    draw.line([(x2, y1), (x2 - length, y1)], fill=color, width=width)
    draw.line([(x2, y1), (x2, y1 + length)], fill=color, width=width)
    draw.line([(x1, y2), (x1 + length, y2)], fill=color, width=width)
    draw.line([(x1, y2), (x1, y2 - length)], fill=color, width=width)
    draw.line([(x2, y2), (x2 - length, y2)], fill=color, width=width)
    draw.line([(x2, y2), (x2 - length, y2)], fill=color, width=width)

def make_phone_mockup(screen_img, target_width=215, target_height=470, accent_color=(0, 255, 136)):
    bezel_h = 8
    bezel_v = 10
    corner_r = 24
    
    phone_w = target_width
    phone_h = target_height
    screen_w = phone_w - (bezel_h * 2)
    screen_h = phone_h - (bezel_v * 2)
    
    screen_resized = screen_img.resize((screen_w, screen_h), Image.Resampling.LANCZOS)
    screen_mask = create_rounded_rect_mask((screen_w, screen_h), radius=18)
    
    phone = Image.new('RGBA', (phone_w, phone_h), (0, 0, 0, 0))
    draw = ImageDraw.Draw(phone)
    
    # Outer frame
    draw.rounded_rectangle(
        [(0, 0), (phone_w - 1, phone_h - 1)],
        radius=corner_r,
        fill=(12, 18, 16, 255),
        outline=(accent_color[0], accent_color[1], accent_color[2], 220),
        width=2
    )
    
    # Inner border
    draw.rounded_rectangle(
        [(2, 2), (phone_w - 3, phone_h - 3)],
        radius=corner_r - 2,
        fill=(6, 10, 8, 255),
        outline=(25, 45, 35, 255),
        width=1
    )
    
    # Paste screen
    phone.paste(screen_resized, (bezel_h, bezel_v), screen_mask)
    
    # Notch
    notch_w = int(phone_w * 0.32)
    notch_h = 9
    notch_x = (phone_w - notch_w) // 2
    notch_y = bezel_v + 3
    draw.rounded_rectangle(
        [(notch_x, notch_y), (notch_x + notch_w, notch_y + notch_h)],
        radius=4,
        fill=(4, 6, 5, 255)
    )
    # Lens
    draw.ellipse(
        [(notch_x + notch_w - 10, notch_y + 2), (notch_x + notch_w - 6, notch_y + 6)],
        fill=(25, 45, 60, 255)
    )
    
    return phone

def make_browser_mockup(screen_img, target_width=710, target_height=365, title="https://command.rakshanet.in", status="LIVE (OK)", font_consola=None, font_bahn=None):
    header_h = 32
    corner_r = 10
    
    screen_w = target_width - 4
    screen_h = target_height - header_h - 2
    
    screen_resized = screen_img.resize((screen_w, screen_h), Image.Resampling.LANCZOS)
    screen_mask = create_rounded_rect_mask((screen_w, screen_h), radius=5)
    
    window = Image.new('RGBA', (target_width, target_height), (0, 0, 0, 0))
    draw = ImageDraw.Draw(window)
    
    # Window chassis
    draw.rounded_rectangle(
        [(0, 0), (target_width - 1, target_height - 1)],
        radius=corner_r,
        fill=(10, 18, 20, 255),
        outline=(0, 229, 255, 220),
        width=2
    )
    
    # Header bar
    header_img = Image.new('RGBA', (target_width, header_h), (14, 26, 28, 255))
    h_draw = ImageDraw.Draw(header_img)
    # Window dots
    h_draw.ellipse([(12, 11), (20, 19)], fill=(255, 95, 86, 255))
    h_draw.ellipse([(26, 11), (34, 19)], fill=(255, 189, 46, 255))
    h_draw.ellipse([(40, 11), (48, 19)], fill=(39, 201, 63, 255))
    
    # URL address bar
    addr_w = int(target_width * 0.46)
    addr_x = (target_width - addr_w) // 2
    h_draw.rounded_rectangle(
        [(addr_x, 6), (addr_x + addr_w, 26)],
        radius=5,
        fill=(20, 36, 40, 255),
        outline=(40, 75, 80, 255),
        width=1
    )
    if font_consola:
        h_draw.text((addr_x + 10, 7), title, font=font_consola, fill=(130, 220, 210, 255))
        
    # Status chip
    status_x = target_width - 120
    h_draw.rounded_rectangle(
        [(status_x, 6), (status_x + 110, 26)],
        radius=5,
        fill=(5, 36, 30, 255),
        outline=(0, 255, 136, 180),
        width=1
    )
    h_draw.ellipse([(status_x + 8, 13), (status_x + 14, 19)], fill=(0, 255, 136, 255))
    if font_bahn:
        h_draw.text((status_x + 18, 7), status, font=font_bahn, fill=(0, 255, 136, 255))
        
    window.paste(header_img, (0, 0))
    draw.line([(0, header_h), (target_width, header_h)], fill=(25, 55, 60, 255), width=1)
    window.paste(screen_resized, (2, header_h + 1), screen_mask)
    
    return window

def build_tag_pill(draw, x, y, text, font, border_color=(0, 255, 136), bg_color=(0, 40, 22), text_color=(0, 255, 136), pad_x=12, pad_y=5):
    bbox = font.getbbox(text)
    tw = bbox[2] - bbox[0]
    th = bbox[3] - bbox[1]
    pw = tw + pad_x * 2
    ph = th + pad_y * 2
    draw.rounded_rectangle([(x, y), (x + pw, y + ph)], radius=5, fill=bg_color, outline=border_color, width=1)
    draw.text((x + pad_x, y + pad_y - 1), text, font=font, fill=text_color)
    return pw, ph

def recreate_slide_5():
    # Pristine base slide
    base_path = r'p:\RakshaNet\slide5_clean_base_150.png'
    slide = Image.open(base_path).convert('RGBA')
    W, H = slide.size
    print(f"Base size: {W}x{H}")
    
    # Fonts
    font_title = ImageFont.truetype(r'C:\Windows\Fonts\bahnschrift.ttf', 32)
    font_sub = ImageFont.truetype(r'C:\Windows\Fonts\bahnschrift.ttf', 17)
    font_label = ImageFont.truetype(r'C:\Windows\Fonts\bahnschrift.ttf', 16)
    font_chip = ImageFont.truetype(r'C:\Windows\Fonts\bahnschrift.ttf', 14)
    font_meta = ImageFont.truetype(r'C:\Windows\Fonts\consola.ttf', 13)
    font_consola = ImageFont.truetype(r'C:\Windows\Fonts\consola.ttf', 11)
    
    # User's 5 images
    m1_img = Image.open(r'C:\Users\Prana\.gemini\antigravity\brain\4f3a9134-7a56-4f1b-a67b-a9fd8aad13be\.user_uploaded\media_1791305412906.jpg') # Live Offline Map
    m2_img = Image.open(r'C:\Users\Prana\.gemini\antigravity\brain\4f3a9134-7a56-4f1b-a67b-a9fd8aad13be\.user_uploaded\media_1791305412909.jpg') # SOS Beacon
    m3_img = Image.open(r'C:\Users\Prana\.gemini\antigravity\brain\4f3a9134-7a56-4f1b-a67b-a9fd8aad13be\.user_uploaded\media_1791305412912.jpg') # Community Chat
    w1_img = Image.open(r'C:\Users\Prana\.gemini\antigravity\brain\4f3a9134-7a56-4f1b-a67b-a9fd8aad13be\.user_uploaded\media_1791305412939.png') # Live GIS Map
    w2_img = Image.open(r'C:\Users\Prana\.gemini\antigravity\brain\4f3a9134-7a56-4f1b-a67b-a9fd8aad13be\.user_uploaded\media_1791305412945.png') # Fleet Dispatch
    
    overlay = Image.new('RGBA', (W, H), (0, 0, 0, 0))
    draw_ov = ImageDraw.Draw(overlay)
    
    # Left 62% spans x=75 to x=1610 (width = 1535 px)
    # y=405 to y=1300 (height = 895 px)
    col_gap = 30
    col1_w = 745  # Mobile Column (fits 3 phones cleanly)
    col2_w = 760  # Web Column (fits 2 desktop screens cleanly)
    
    col1_x = 75
    col2_x = col1_x + col1_w + col_gap
    col_y = 405
    col_h = 895
    
    # ==========================================
    # PILLAR 1: MOBILE APPLICATION (APK)
    # ==========================================
    draw_ov.rounded_rectangle(
        [(col1_x, col_y), (col1_x + col1_w, col_y + col_h)],
        radius=16,
        fill=(4, 18, 12, 235),
        outline=(0, 255, 136, 180),
        width=2
    )
    draw_hud_corners(draw_ov, (col1_x, col_y, col1_x + col1_w, col_y + col_h), (0, 255, 136, 255), length=22, width=3)
    draw_ov.line([(col1_x + 35, col_y + 2), (col1_x + col1_w - 35, col_y + 2)], fill=(0, 255, 136, 255), width=3)
    
    # Header
    build_tag_pill(draw_ov, col1_x + 25, col_y + 18, "FIELD RESPONDER SUITE", font_meta, (0, 255, 136), (0, 45, 25), (0, 255, 136))
    draw_ov.text((col1_x + col1_w - 180, col_y + 20), "[ ANDROID APK ]", font=font_meta, fill=(0, 255, 136, 180))
    draw_ov.text((col1_x + 25, col_y + 48), "Mobile Application", font=font_title, fill=(255, 255, 255, 255))
    draw_ov.text((col1_x + 25, col_y + 88), "Offline Mesh P2P", font=font_sub, fill=(120, 220, 175, 255))
    
    # 3 Mobile Phones side-by-side!
    phone_w = 215
    phone_h = 475
    phone_gap = 20
    p_start_x = col1_x + (col1_w - (3 * phone_w + 2 * phone_gap)) // 2
    phone_y = col_y + 125
    
    px1 = p_start_x
    px2 = px1 + phone_w + phone_gap
    px3 = px2 + phone_w + phone_gap
    
    phone1 = make_phone_mockup(m1_img, target_width=phone_w, target_height=phone_h, accent_color=(0, 229, 255))
    phone2 = make_phone_mockup(m2_img, target_width=phone_w, target_height=phone_h, accent_color=(255, 80, 80))
    phone3 = make_phone_mockup(m3_img, target_width=phone_w, target_height=phone_h, accent_color=(255, 210, 31))
    
    overlay.paste(phone1, (px1, phone_y), phone1)
    overlay.paste(phone2, (px2, phone_y), phone2)
    overlay.paste(phone3, (px3, phone_y), phone3)
    
    # 2-3 Word Labels under each phone
    l_y = phone_y + phone_h + 10
    
    # Label 1: Live Offline Map (3 words)
    build_tag_pill(draw_ov, px1 + (phone_w - 150) // 2, l_y, "Live Offline Map", font_chip, (0, 229, 255), (10, 30, 35), (0, 229, 255), pad_x=14, pad_y=6)
    
    # Label 2: Emergency SOS Beacon (3 words)
    build_tag_pill(draw_ov, px2 + (phone_w - 180) // 2, l_y, "Emergency SOS Beacon", font_chip, (255, 80, 80), (40, 15, 15), (255, 120, 120), pad_x=14, pad_y=6)
    
    # Label 3: Community Mesh Chat (3 words)
    build_tag_pill(draw_ov, px3 + (phone_w - 180) // 2, l_y, "Community Mesh Chat", font_chip, (255, 210, 31), (40, 35, 10), (255, 210, 31), pad_x=14, pad_y=6)
    
    # 3 Bottom Highlight Chips (2-3 words each)
    chip_y = l_y + 48
    c_w = (col1_w - 50 - 20) // 3
    
    build_tag_pill(draw_ov, col1_x + 25, chip_y, "Zero-Internet Mesh", font_chip, (0, 255, 136), (6, 28, 18), (220, 255, 235), pad_x=16, pad_y=8)
    build_tag_pill(draw_ov, col1_x + 25 + c_w + 10, chip_y, "GPS-Free Radar", font_chip, (0, 229, 255), (6, 26, 32), (210, 245, 255), pad_x=22, pad_y=8)
    build_tag_pill(draw_ov, col1_x + 25 + (c_w + 10) * 2, chip_y, "On-Device AI Help", font_chip, (255, 210, 31), (32, 28, 8), (255, 240, 200), pad_x=18, pad_y=8)
    
    # ==========================================
    # PILLAR 2: WEB COMMAND CENTER
    # ==========================================
    draw_ov.rounded_rectangle(
        [(col2_x, col_y), (col2_x + col2_w, col_y + col_h)],
        radius=16,
        fill=(6, 16, 22, 235),
        outline=(0, 229, 255, 180),
        width=2
    )
    draw_hud_corners(draw_ov, (col2_x, col_y, col2_x + col2_w, col_y + col_h), (0, 229, 255, 255), length=22, width=3)
    draw_ov.line([(col2_x + 35, col_y + 2), (col2_x + col2_w - 35, col_y + 2)], fill=(0, 229, 255, 255), width=3)
    
    # Header
    build_tag_pill(draw_ov, col2_x + 25, col_y + 18, "TACTICAL COMMAND HQ", font_meta, (0, 229, 255), (0, 35, 45), (0, 229, 255))
    draw_ov.text((col2_x + col2_w - 180, col_y + 20), "[ WEB CONSOLE ]", font=font_meta, fill=(0, 229, 255, 180))
    draw_ov.text((col2_x + 25, col_y + 48), "Command Center", font=font_title, fill=(255, 255, 255, 255))
    draw_ov.text((col2_x + 25, col_y + 88), "Tactical GIS & Dispatch", font=font_sub, fill=(120, 215, 235, 255))
    
    # 2 Desktop Windows stacked vertically!
    browser_w = col2_w - 50  # 710 px
    browser_h = 320          # 320 px
    bx = col2_x + 25
    
    # Window 1: Tactical Live GIS (Media 4)
    by1 = col_y + 125
    b1_mockup = make_browser_mockup(
        w1_img,
        target_width=browser_w,
        target_height=browser_h,
        title="https://command.rakshanet.in/live-map",
        status="GIS (ONLINE)",
        font_consola=font_consola,
        font_bahn=font_meta
    )
    overlay.paste(b1_mockup, (bx, by1), b1_mockup)
    
    # 2-3 Word Label under Window 1
    build_tag_pill(draw_ov, bx + 15, by1 + browser_h + 8, "Tactical Live GIS", font_chip, (0, 229, 255), (10, 28, 36), (0, 229, 255), pad_x=14, pad_y=5)
    build_tag_pill(draw_ov, bx + 200, by1 + browser_h + 8, "Geospatial Clusters", font_chip, (0, 255, 136), (6, 28, 18), (180, 255, 220), pad_x=14, pad_y=5)
    
    # Window 2: Fleet Dispatch & Rescue Ops (Media 5)
    by2 = by1 + browser_h + 52
    b2_mockup = make_browser_mockup(
        w2_img,
        target_width=browser_w,
        target_height=browser_h,
        title="https://command.rakshanet.in/rescue-ops",
        status="MILP (READY)",
        font_consola=font_consola,
        font_bahn=font_meta
    )
    overlay.paste(b2_mockup, (bx, by2), b2_mockup)
    
    # 2-3 Word Label under Window 2
    build_tag_pill(draw_ov, bx + 15, by2 + browser_h + 8, "Fleet Dispatch Optimizer", font_chip, (255, 210, 31), (32, 28, 8), (255, 210, 31), pad_x=14, pad_y=5)
    build_tag_pill(draw_ov, bx + 245, by2 + browser_h + 8, "Optimal Unit Routing", font_chip, (0, 229, 255), (10, 28, 36), (200, 245, 255), pad_x=14, pad_y=5)
    
    # Composite overlay onto clean base slide
    final_slide = Image.alpha_composite(slide, overlay)
    
    # Save outputs:
    # 1. High-Res JPEG
    jpg_path = r'p:\RakshaNet\Slide_5_Key_Features_UIUX.jpg'
    final_slide.convert('RGB').save(jpg_path, 'JPEG', quality=95)
    print(f"Saved JPEG: {jpg_path}")
    
    # 2. PNG slide in workspace
    png_path = r'p:\RakshaNet\slide_page_5_ui_ux.png'
    final_slide.convert('RGB').save(png_path, 'PNG', quality=95)
    print(f"Saved PNG: {png_path}")
    
    # 3. Brain artifact
    brain_path = r'C:\Users\Prana\.gemini\antigravity\brain\4f3a9134-7a56-4f1b-a67b-a9fd8aad13be\pdf_slides\page_5.png'
    final_slide.convert('RGB').save(brain_path, 'PNG', quality=95)
    print(f"Saved Brain slide: {brain_path}")
    
    # 4. Generate PPTX format!
    pptx_path = r'p:\RakshaNet\Slide_5_Key_Features_UIUX.pptx'
    prs = Presentation()
    # 16:9 widescreen
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank_layout = prs.slide_layouts[6] # blank layout
    ppt_slide = prs.slides.add_slide(blank_layout)
    # Add slide image filling the entire slide
    ppt_slide.shapes.add_picture(jpg_path, Inches(0), Inches(0), Inches(13.333), Inches(7.5))
    prs.save(pptx_path)
    print(f"Saved PPTX: {pptx_path}")
    
    # 5. Also update presentation PDF
    pdf_in = r'p:\RakshaNet\template_backup.pdf'
    pdf_out = r'p:\RakshaNet\RakshaNet_CODE_APEX_Presentation.pdf'
    if os.path.exists(pdf_in):
        doc = pymupdf.open(pdf_in)
        page = doc[4]
        rect = page.rect
        page.insert_image(rect, filename=png_path)
        doc.save(pdf_out)
        doc.close()
        print(f"Saved PDF: {pdf_out}")

if __name__ == "__main__":
    recreate_slide_5()
